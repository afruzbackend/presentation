"""Mavzu matnidan taqdimot: PowerPoint (.pptx) va PDF — ikkalasi bir xil ko'rinishda, Samarqand uslubida.

    slides = outline(topic)          # slaydlar mazmuni (sarlavha, reja, maqsad, har bir kichik mavzu, atamalar...)
    pptx_bytes = render_pptx(topic)  # PowerPoint'da ko'rsatish / tahrirlash uchun
    pdf_bytes = render_pdf(topic)    # saytning o'zida (brauzerda) ko'rinadi

Matn joylashuvi (qator bo'linishi, shrift o'lchami) bir marta hisoblanadi — ikkala format bir xil chiqadi.
Shriftlar: Georgia (sarlavhalar) va Segoe UI (matn) — Windows'da bor.
"""
import html
import io
import re
from functools import lru_cache
from pathlib import Path

from PIL import ImageFont

# ------------------------------------------------------------------ ranglar, o'lchamlar, shriftlar
W, H = 13.333, 7.5  # dyuym, 16:9
PAPER, PAPER2, CARD = '#F4EBDD', '#ECDFCB', '#FBF6EC'
INK, INK2, MUTED = '#1B2A4A', '#3A4560', '#6B6457'
TERRA, OLTIN, LINE = '#C2562D', '#B9862A', '#D8C8AE'

FONT_DIR = Path('C:/Windows/Fonts')
FONTS = {
    ('serif', False): 'georgia.ttf', ('serif', True): 'georgiab.ttf',
    ('sans', False): 'segoeui.ttf', ('sans', True): 'segoeuib.ttf',
}
FACE = {'serif': 'Georgia', 'sans': 'Segoe UI'}


class FontMissing(Exception):
    pass


@lru_cache(maxsize=None)
def _pil(family, bold, size):
    path = FONT_DIR / FONTS[(family, bold)]
    if not path.exists():
        raise FontMissing(f'Shrift topilmadi: {path}')
    return ImageFont.truetype(str(path), size=200), size  # 200 px da o'lchab, keyin masshtablaymiz


def text_width(text, family, bold, size):
    """Matn kengligi, dyuymda."""
    font, _ = _pil(family, bold, 0)
    return font.getlength(text) * size / 200 / 72


# ------------------------------------------------------------------ matnni tozalash va qisqartirish
def plain(s):
    s = html.unescape(re.sub(r'<[^>]+>', '', s or ''))
    return re.sub(r'\s+', ' ', s).strip()


def clip(s, n):
    """Gap yoki so'z chegarasida qisqartirish."""
    s = plain(s)
    if len(s) <= n:
        return s
    cut = s[:n]
    m = max(cut.rfind('. '), cut.rfind('; '))
    if m > n * 0.55:
        return cut[:m + 1]
    return cut[:cut.rfind(' ')].rstrip(',;:—–- ') + '…'


def first_sentences(s, n=230):
    s = plain(s)
    out = ''
    for m in re.finditer(r'.+?[.!?](?=\s|$)', s):
        if out and len(out) + len(m.group(0)) > n:
            break
        out += m.group(0)
    return clip(out or s, n)


def split_item(item):
    """'<b>Kalit:</b> izoh' -> ('Kalit', 'izoh')"""
    m = re.match(r'^\s*<b>(.+?)</b>\s*[:—–-]?\s*(.*)$', item or '', re.S)
    if m:
        return plain(m.group(1)).rstrip(': '), plain(m.group(2))
    p = plain(item)
    i = p.find(':')
    if 2 < i < 46:
        return p[:i], p[i + 1:].strip()
    return '', p


# ------------------------------------------------------------------ slaydlar mazmuni
def outline(topic):
    course = topic.course
    slides = [{'t': 'title'}]
    sections = list(topic.sections.all())
    if sections:
        slides.append({'t': 'plan', 'title': 'Mashg‘ulot rejasi', 'items': [s.title for s in sections]})
    if topic.goal or topic.result:
        slides.append({'t': 'goal', 'title': 'Maqsad va kutilayotgan natija',
                       'goal': clip(topic.goal, 420), 'result': clip(topic.result, 420)})
    for i, sec in enumerate(sections, 1):
        eyebrow = f'{i}-reja · {sec.title}'
        for g in _groups(sec):
            slides.append({**g, 'eyebrow': eyebrow, 'title': g['title'] or sec.title})
    terms = [(g.term, clip(g.definition, 150)) for g in topic.terms.all()[:6]]
    if terms:
        slides.append({'t': 'terms', 'title': 'Asosiy atamalar', 'items': terms})
    qs = [plain(q) for q in (topic.questions or [])][:7]
    if qs:
        slides.append({'t': 'questions', 'title': 'Mustahkamlash uchun savollar', 'items': qs})
    slides.append({'t': 'end', 'author': course.author})
    return slides


def _groups(sec):
    """Bo'limni kichik sarlavhalar bo'yicha slaydlarga ajratish: sarlavha + matn + ro'yxat yoki jadval."""
    groups, cur = [], None
    for b in sec.blocks:
        if b['t'] == 'h' or cur is None:
            cur = {'title': plain(b.get('text', '')) if b['t'] == 'h' else '', 'ps': [], 'ul': [], 'table': None}
            groups.append(cur)
            if b['t'] == 'h':
                continue
        if b['t'] == 'p':
            cur['ps'].append(b['html'])
        elif b['t'] == 'ul':
            cur['ul'] += b['items']
        elif b['t'] == 'table' and cur['table'] is None:
            cur['table'] = b['rows']
    out = []
    for g in groups:
        lead = first_sentences(g['ps'][0], 260) if g['ps'] else ''
        if g['table']:
            rows = [[clip(c, 70) for c in r[:4]] for r in g['table'][:6]]
            out.append({'t': 'table', 'title': g['title'], 'lead': lead, 'rows': rows})
        elif g['ul']:
            items = [split_item(x) for x in g['ul'][:6]]
            out.append({'t': 'bullets', 'title': g['title'], 'lead': lead,
                        'items': [(k, clip(v, 150)) for k, v in items]})
        elif g['ps']:
            rest = [('', first_sentences(p, 170)) for p in g['ps'][1:5]]
            out.append({'t': 'bullets', 'title': g['title'], 'lead': lead, 'items': rest})
    return [g for g in out if g['lead'] or g.get('items') or g.get('rows')]


# ------------------------------------------------------------------ joylashuv: chizish buyruqlari
def wrap(runs, width, size):
    """runs=[(matn, qalin)] -> qatorlar [[(matn, qalin)], ...] (berilgan kenglikka sig'adigan)."""
    lines, line, x = [], [], 0.0
    for text, bold in runs:
        for word in re.findall(r'\S+\s*', text):
            w = text_width(word.rstrip(), 'sans', bold, size)
            if line and x + w > width:
                lines.append(line)
                line, x = [], 0.0
            line.append((word, bold))
            x += text_width(word, 'sans', bold, size)
    if line:
        lines.append(line)
    return lines


def wrap_family(text, width, size, family, bold):
    lines, line = [], ''
    for word in text.split():
        test = f'{line} {word}'.strip()
        if line and text_width(test, family, bold, size) > width:
            lines.append(line)
            line = word
        else:
            line = test
    if line:
        lines.append(line)
    return lines


class Canvas:
    """Formatdan mustaqil chizish buyruqlari ro'yxati (keyin PPTX yoki PDF'ga aylantiriladi)."""

    def __init__(self):
        self.ops = []

    def rect(self, x, y, w, h, fill, line=None, radius=0):
        self.ops.append(('rect', x, y, w, h, fill, line, radius))

    def star(self, cx, cy, r, fill):
        self.ops.append(('star', cx, cy, r, fill))

    def diamond(self, cx, cy, r, fill):
        self.ops.append(('diamond', cx, cy, r, fill))

    def arch(self, x, y, w, h, fill):
        self.ops.append(('arch', x, y, w, h, fill))

    def text(self, x, y, w, lines, size, family='sans', color=INK, lh=1.25, align='l'):
        """lines: [[(matn, qalin)], ...] — qatorlar oldindan bo'lingan."""
        self.ops.append(('text', x, y, w, lines, size, family, color, lh, align))
        return y + len(lines) * size * lh / 72

    def table(self, x, y, col_w, rows, size, accent):
        self.ops.append(('table', x, y, col_w, rows, size, accent))


def _frame(c, topic, page, total, accent):
    c.rect(0, 0, W, H, PAPER)
    c.rect(0, 0, W, 0.14, accent)
    for i in range(27):  # koshin hoshiyasi
        c.diamond(0.25 + i * 0.5, 0.3, 0.07, OLTIN if i % 2 else accent)
    c.rect(0.6, H - 0.62, W - 1.2, 0.012, LINE)
    foot = [[(f'{topic.course.title} · {topic.number}-mavzu', False)]]
    c.text(0.6, H - 0.5, 8, foot, 10, color=MUTED)
    c.text(W - 2.6, H - 0.5, 2, [[(f'{page} / {total}', True)]], 10, color=MUTED, align='r')


def _heading(c, eyebrow, title, accent):
    y = 0.62
    if eyebrow:
        c.text(0.8, y, W - 1.6, [[(clip(eyebrow, 95).upper(), True)]], 11, color=accent)
        y += 0.32
    size = 30
    lines = wrap_family(title, W - 1.6, size, 'serif', True)
    while len(lines) > 2 and size > 20:
        size -= 2
        lines = wrap_family(title, W - 1.6, size, 'serif', True)
    y = c.text(0.8, y, W - 1.6, [[(ln, True)] for ln in lines[:2]], size, family='serif', color=INK, lh=1.15)
    c.rect(0.8, y + 0.1, 0.9, 0.05, accent)
    return y + 0.38


def _fit_items(items, width, top, bottom, lead, sizes=(26, 24, 22, 21, 20, 19, 18, 17, 16, 15, 14, 13)):
    """Matn va ro'yxat sig'adigan eng katta shrift o'lchamini topish."""
    for size in sizes:
        lead_lines = wrap([(lead, False)], width, size + 1) if lead else []
        y = top + (len(lead_lines) * (size + 1) * 1.3 / 72 + 0.22 if lead else 0)
        wrapped = []
        for k, v in items:
            runs = [(k + (': ' if v else ''), True)] if k else []
            runs.append((v, False))
            ls = wrap(runs, width - 0.45, size)
            wrapped.append(ls)
            y += len(ls) * size * 1.28 / 72 + 0.16
        if y <= bottom:
            return size, lead_lines, wrapped
    return size, lead_lines, wrapped


def layout(topic):
    """Har bir slayd uchun Canvas (chizish buyruqlari)."""
    accent = topic.color or TERRA
    slides = outline(topic)
    pages = []
    for n, s in enumerate(slides, 1):
        c = Canvas()
        _frame(c, topic, n, len(slides), accent)
        getattr(_SLIDES, s['t'])(c, topic, s, accent)
        pages.append(c)
    return pages


class _SLIDES:
    @staticmethod
    def title(c, topic, s, accent):
        c.arch(0.8, 0.85, 3.6, 5.6, accent)
        c.star(2.6, 2.75, 1.05, PAPER)
        c.text(1.4, 2.35, 2.4, [[(str(topic.number), True)]], 46, family='serif', color=accent, align='c', lh=1.0)
        c.text(1.0, 4.5, 3.2, [[('MAVZU', True)]], 14, color=PAPER, align='c')
        c.text(1.0, 4.85, 3.2, [[(f'{topic.kind} · {topic.hours} soat', False)]], 13, color=PAPER, align='c')
        x, w = 5.0, W - 5.8
        y = c.text(x, 1.2, w, [[(f'{topic.course.title.upper()}', True)]], 13, color=accent)
        size = 40
        lines = wrap_family(topic.title, w, size, 'serif', True)
        while len(lines) > 4 and size > 26:
            size -= 2
            lines = wrap_family(topic.title, w, size, 'serif', True)
        y = c.text(x, y + 0.25, w, [[(ln, True)] for ln in lines], size, family='serif', color=INK, lh=1.12)
        c.rect(x, y + 0.25, 1.2, 0.06, accent)
        meta = [topic.course.author, topic.course.author_role, topic.course.college]
        y += 0.6
        for i, m in enumerate(x for x in meta if x):
            y = c.text(x, y, w, [[(clip(m, 90), i == 0)]], 20 if i == 0 else 15, color=INK if i == 0 else MUTED) + 0.08

    @staticmethod
    def plan(c, topic, s, accent):
        y = _heading(c, f'{topic.number}-mavzu', s['title'], accent)
        items = s['items'][:6]
        cols = 2 if len(items) > 3 else 1
        cw = (W - 1.6 - 0.4 * (cols - 1)) / cols
        rows = -(-len(items) // cols)
        ch = min(1.35, (H - 0.9 - y - 0.25 * (rows - 1)) / rows)
        for i, t in enumerate(items):
            col, row = (i // rows, i % rows) if cols == 2 else (0, i)
            x, yy = 0.8 + col * (cw + 0.4), y + row * (ch + 0.25)
            c.rect(x, yy, cw, ch, CARD, LINE, radius=0.08)
            c.star(x + 0.6, yy + ch / 2, 0.36, accent)
            c.text(x + 0.3, yy + ch / 2 - 0.2, 0.6, [[(str(i + 1), True)]], 20, family='serif', color=PAPER, align='c', lh=1.0)
            size = 24
            ls = wrap_family(t, cw - 1.5, size, 'serif', True)
            while len(ls) * size * 1.2 / 72 > ch - 0.2 and size > 12:
                size -= 1
                ls = wrap_family(t, cw - 1.5, size, 'serif', True)
            th = len(ls) * size * 1.2 / 72
            c.text(x + 1.2, yy + (ch - th) / 2, cw - 1.4, [[(ln, True)] for ln in ls], size, family='serif', lh=1.2)

    @staticmethod
    def goal(c, topic, s, accent):
        y = _heading(c, f'{topic.number}-mavzu', s['title'], accent)
        boxes = [(lbl, txt) for lbl, txt in (('Maqsad', s['goal']), ('Kutilayotgan natija', s['result'])) if txt]
        bw = (W - 1.6 - 0.4 * (len(boxes) - 1)) / len(boxes)
        bh = H - 0.95 - y
        for i, (lbl, txt) in enumerate(boxes):
            x = 0.8 + i * (bw + 0.4)
            c.rect(x, y, bw, bh, CARD, LINE, radius=0.08)
            c.rect(x, y, bw, 0.09, accent if i == 0 else OLTIN)
            c.text(x + 0.35, y + 0.35, bw - 0.7, [[(lbl.upper(), True)]], 13, color=accent if i == 0 else OLTIN)
            size = 26
            ls = wrap([(txt, False)], bw - 0.7, size)
            while len(ls) * size * 1.35 / 72 > bh - 1.1 and size > 13:
                size -= 1
                ls = wrap([(txt, False)], bw - 0.7, size)
            c.text(x + 0.35, y + 0.8, bw - 0.7, ls, size, color=INK2, lh=1.35)

    @staticmethod
    def bullets(c, topic, s, accent):
        y = _heading(c, s['eyebrow'], s['title'], accent)
        width, bottom = W - 1.6, H - 0.85
        size, lead_lines, wrapped = _fit_items(s['items'], width, y, bottom, s['lead'])
        if lead_lines:
            c.rect(0.8, y - 0.04, 0.05, len(lead_lines) * (size + 1) * 1.3 / 72 + 0.08, OLTIN)
            y = c.text(1.05, y, width - 0.25, lead_lines, size + 1, color=INK2, lh=1.3) + 0.22
        for ls in wrapped:
            if y + size * 1.28 / 72 > bottom:
                break
            c.diamond(0.98, y + size * 0.62 / 72, 0.085, accent)
            y = c.text(1.25, y, width - 0.45, ls, size, color=INK, lh=1.28) + 0.16

    @staticmethod
    def table(c, topic, s, accent):
        y = _heading(c, s['eyebrow'], s['title'], accent)
        width, bottom = W - 1.6, H - 0.85
        if s['lead']:
            ls = wrap([(s['lead'], False)], width - 0.25, 19)[:3]
            c.rect(0.8, y - 0.04, 0.05, len(ls) * 19 * 1.3 / 72 + 0.08, OLTIN)
            y = c.text(1.05, y, width - 0.25, ls, 19, color=INK2, lh=1.3) + 0.25
        rows = s['rows']
        ncol = max(len(r) for r in rows)
        rows = [r + [''] * (ncol - len(r)) for r in rows]
        # ustun kengligi matn uzunligiga qarab
        weights = [max(8, max(len(r[i]) for r in rows)) for i in range(ncol)]
        col_w = [width * w / sum(weights) for w in weights]
        for size in (19, 18, 17, 16, 15, 14, 13, 12, 11):
            cells = [[wrap_family(txt, cw - 0.24, size, 'sans', ri == 0) for txt, cw in zip(r, col_w)] for ri, r in enumerate(rows)]
            heights = [max(len(x) for x in r) * size * 1.25 / 72 + 0.2 for r in cells]
            if y + sum(heights) <= bottom:
                break
        while rows and y + sum(heights) > bottom:  # baribir sig'masa — oxirgi qatorlar tushib qoladi
            rows, cells, heights = rows[:-1], cells[:-1], heights[:-1]
        c.table(0.8, y, col_w, [(cl, h) for cl, h in zip(cells, heights)], size, accent)

    @staticmethod
    def terms(c, topic, s, accent):
        y = _heading(c, f'{topic.number}-mavzu', s['title'], accent)
        items = s['items']
        cols = 2 if len(items) > 2 else 1
        rows = -(-len(items) // cols)
        cw = (W - 1.6 - 0.35 * (cols - 1)) / cols
        ch = (H - 0.9 - y - 0.25 * (rows - 1)) / rows
        for i, (term, d) in enumerate(items):
            x, yy = 0.8 + (i % cols) * (cw + 0.35), y + (i // cols) * (ch + 0.25)
            c.rect(x, yy, cw, ch, CARD, LINE, radius=0.08)
            c.rect(x, yy, 0.08, ch, accent)
            size = 19
            while True:
                tl = wrap_family(term, cw - 0.55, size + 3, 'serif', True)[:2]
                dl = wrap([(d, False)], cw - 0.55, size)
                need = len(tl) * (size + 3) * 1.15 / 72 + 0.1 + len(dl) * size * 1.3 / 72
                if need <= ch - 0.3 or size <= 11:
                    break
                size -= 1
            yy2 = c.text(x + 0.3, yy + 0.18, cw - 0.5, [[(ln, True)] for ln in tl], size + 3, family='serif', lh=1.15) + 0.1
            max_lines = max(1, int((yy + ch - 0.12 - yy2) / (size * 1.3 / 72)))
            c.text(x + 0.3, yy2, cw - 0.5, dl[:max_lines], size, color=INK2, lh=1.3)

    @staticmethod
    def questions(c, topic, s, accent):
        y = _heading(c, f'{topic.number}-mavzu', s['title'], accent)
        bottom = H - 0.85
        for size in (26, 24, 22, 20, 19, 18, 17, 16, 15, 14):
            wrapped = [wrap_family(q, W - 2.6, size, 'sans', False) for q in s['items']]
            need = sum(len(w) * size * 1.3 / 72 + 0.2 for w in wrapped)
            if y + need <= bottom:
                break
        for i, ls in enumerate(wrapped, 1):
            if y + size * 1.3 / 72 > bottom:
                break
            c.star(1.08, y + size * 0.66 / 72, 0.22, accent)
            c.text(0.86, y + size * 0.66 / 72 - 0.12, 0.44, [[(str(i), True)]], 12, color=PAPER, align='c', lh=1.0)
            y = c.text(1.55, y, W - 2.6, [[(ln, False)] for ln in ls], size, lh=1.3) + 0.2

    @staticmethod
    def end(c, topic, s, accent):
        c.star(W / 2, 2.45, 0.95, accent)
        c.star(W / 2, 2.45, 0.55, PAPER)
        c.text(0.8, 3.75, W - 1.6, [[('E’tiboringiz uchun rahmat!', True)]], 40, family='serif', align='c')
        c.text(0.8, 4.65, W - 1.6, [[('Savollaringiz bo‘lsa, marhamat.', False)]], 20, color=INK2, align='c')
        c.text(0.8, 5.35, W - 1.6, [[(s['author'], True)]], 16, color=accent, align='c')


# ------------------------------------------------------------------ PowerPoint
def render_pptx(topic):
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
    from pptx.util import Emu, Inches, Pt

    def rgb(h):
        return RGBColor.from_string(h.lstrip('#'))

    def fill(shape, color, line=None):
        shape.fill.solid()
        shape.fill.fore_color.rgb = rgb(color)
        if line:
            shape.line.color.rgb = rgb(line)
            shape.line.width = Pt(1)
        else:
            shape.line.fill.background()
        shape.shadow.inherit = False

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(W), Inches(H)
    blank = prs.slide_layouts[6]
    prs.core_properties.title = f'{topic.number}-mavzu. {topic.title}'
    prs.core_properties.author = topic.course.author

    for page in layout(topic):
        sl = prs.slides.add_slide(blank)
        sh = sl.shapes
        for op in page.ops:
            kind = op[0]
            if kind == 'rect':
                _, x, y, w, h, color, line, radius = op
                shp = sh.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
                                   Inches(x), Inches(y), Inches(w), Inches(h))
                if radius:
                    shp.adjustments[0] = min(0.5, radius / min(w, h))
                fill(shp, color, line)
            elif kind in ('star', 'diamond'):
                _, cx, cy, r, color = op
                side = r * 2 / 1.414 if kind == 'star' else r * 1.414
                rots = (0, 45) if kind == 'star' else (45,)
                for rot in rots:  # sakkiz qirrali yulduz = ikki kvadrat (biri 45° burilgan)
                    shp = sh.add_shape(MSO_SHAPE.RECTANGLE, Inches(cx - side / 2), Inches(cy - side / 2), Inches(side), Inches(side))
                    shp.rotation = rot
                    fill(shp, color)
            elif kind == 'arch':
                _, x, y, w, h, color = op
                shp = sh.add_shape(MSO_SHAPE.ROUND_2_SAME_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
                shp.adjustments[0] = 0.5  # yuqori burchaklar yarim aylana — ravoq
                shp.adjustments[1] = 0.0
                fill(shp, color)
            elif kind == 'text':
                _, x, y, w, lines, size, family, color, lh, align = op
                hgt = len(lines) * size * lh / 72 + 0.05
                tb = sh.add_textbox(Inches(x), Inches(y), Inches(w * 1.04), Inches(hgt))
                tf = tb.text_frame
                tf.word_wrap = True
                tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = Emu(0)
                tf.vertical_anchor = MSO_ANCHOR.TOP
                p = tf.paragraphs[0]
                p.line_spacing = Pt(size * lh)  # aniq oraliq — PDF bilan bir xil
                p.alignment = {'l': PP_ALIGN.LEFT, 'c': PP_ALIGN.CENTER, 'r': PP_ALIGN.RIGHT}[align]
                for i, ln in enumerate(lines):  # qatorlar PDF'dagidek aynan shu joyda bo'linadi
                    if i:
                        p.add_line_break()
                    merged = []
                    for t, b in ln:
                        if merged and merged[-1][1] == b:
                            merged[-1][0] += t
                        else:
                            merged.append([t, b])
                    merged[-1][0] = merged[-1][0].rstrip()
                    for t, b in merged:
                        r = p.add_run()
                        r.text = t
                        r.font.size = Pt(size)
                        r.font.bold = b
                        r.font.name = FACE[family]
                        r.font.color.rgb = rgb(color)
            elif kind == 'table':
                _, x, y, col_w, rows, size, accent = op
                total_h = sum(h for _, h in rows)
                gt = sh.add_table(len(rows), len(col_w), Inches(x), Inches(y), Inches(sum(col_w)), Inches(total_h)).table
                for i, cw in enumerate(col_w):
                    gt.columns[i].width = Inches(cw)
                for ri, (cells, h) in enumerate(rows):
                    gt.rows[ri].height = Inches(h)
                    for ci, ls in enumerate(cells):
                        cell = gt.cell(ri, ci)
                        cell.fill.solid()
                        cell.fill.fore_color.rgb = rgb(accent if ri == 0 else (CARD if ri % 2 else PAPER2))
                        cell.margin_left = cell.margin_right = Inches(0.12)
                        cell.margin_top = cell.margin_bottom = Inches(0.08)
                        tf = cell.text_frame
                        tf.word_wrap = True
                        p = tf.paragraphs[0]
                        r = p.add_run()
                        r.text = ' '.join(ls)
                        r.font.size = Pt(size)
                        r.font.bold = ri == 0
                        r.font.name = FACE['sans']
                        r.font.color.rgb = rgb(PAPER if ri == 0 else INK)
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


# ------------------------------------------------------------------ PDF
@lru_cache(maxsize=None)
def _pdf_fonts():
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    names = {}
    for (family, bold), file in FONTS.items():
        name = f'{FACE[family].replace(" ", "")}{"-Bold" if bold else ""}'
        path = FONT_DIR / file
        if not path.exists():
            raise FontMissing(f'Shrift topilmadi: {path}')
        pdfmetrics.registerFont(TTFont(name, str(path)))
        names[(family, bold)] = name
    return names


def render_pdf(topic):
    from reportlab.lib.colors import HexColor
    from reportlab.pdfgen import canvas

    fonts = _pdf_fonts()
    PT = 72
    buf = io.BytesIO()
    pdf = canvas.Canvas(buf, pagesize=(W * PT, H * PT))
    pdf.setTitle(f'{topic.number}-mavzu. {topic.title}')
    pdf.setAuthor(topic.course.author)

    def Y(y):
        return (H - y) * PT

    def poly(points, color):
        pdf.setFillColor(HexColor(color))
        p = pdf.beginPath()
        p.moveTo(*points[0])
        for pt in points[1:]:
            p.lineTo(*pt)
        p.close()
        pdf.drawPath(p, fill=1, stroke=0)

    def square(cx, cy, half, rot, color):
        import math
        a = math.radians(rot)
        pts = []
        for dx, dy in ((-half, -half), (half, -half), (half, half), (-half, half)):
            pts.append(((cx + dx * math.cos(a) - dy * math.sin(a)) * PT, Y(cy + dx * math.sin(a) + dy * math.cos(a))))
        poly(pts, color)

    for page in layout(topic):
        for op in page.ops:
            kind = op[0]
            if kind == 'rect':
                _, x, y, w, h, color, line, radius = op
                pdf.setFillColor(HexColor(color))
                if line:
                    pdf.setStrokeColor(HexColor(line))
                    pdf.setLineWidth(1)
                if radius:
                    pdf.roundRect(x * PT, Y(y + h), w * PT, h * PT, radius * PT, fill=1, stroke=1 if line else 0)
                else:
                    pdf.rect(x * PT, Y(y + h), w * PT, h * PT, fill=1, stroke=1 if line else 0)
            elif kind == 'star':
                _, cx, cy, r, color = op
                for rot in (0, 45):
                    square(cx, cy, r / 1.414, rot, color)
            elif kind == 'diamond':
                _, cx, cy, r, color = op
                square(cx, cy, r / 1.414, 45, color)
            elif kind == 'arch':
                _, x, y, w, h, color = op
                pdf.setFillColor(HexColor(color))
                rad = w / 2
                p = pdf.beginPath()
                p.moveTo(x * PT, Y(y + h))
                p.lineTo(x * PT, Y(y + rad))
                p.arcTo(x * PT, Y(y + 2 * rad), (x + w) * PT, Y(y), startAng=180, extent=-180)
                p.lineTo((x + w) * PT, Y(y + h))
                p.close()
                pdf.drawPath(p, fill=1, stroke=0)
            elif kind == 'text':
                _, x, y, w, lines, size, family, color, lh, align = op
                pdf.setFillColor(HexColor(color))
                for i, ln in enumerate(lines):
                    base = Y(y + (i * lh + 0.92) * size / 72)
                    lw = sum(text_width(t, family, b, size) for t, b in ln) - (
                        text_width(ln[-1][0], family, ln[-1][1], size) - text_width(ln[-1][0].rstrip(), family, ln[-1][1], size))
                    cx = x if align == 'l' else (x + (w - lw) / 2 if align == 'c' else x + w - lw)
                    for t, b in ln:
                        pdf.setFont(fonts[(family, b)], size)
                        pdf.drawString(cx * PT, base, t)
                        cx += text_width(t, family, b, size)
            elif kind == 'table':
                _, x, y, col_w, rows, size, accent = op
                yy = y
                for ri, (cells, h) in enumerate(rows):
                    xx = x
                    for ci, ls in enumerate(cells):
                        bg = accent if ri == 0 else (CARD if ri % 2 else PAPER2)
                        pdf.setFillColor(HexColor(bg))
                        pdf.setStrokeColor(HexColor(PAPER))
                        pdf.setLineWidth(1.2)
                        pdf.rect(xx * PT, Y(yy + h), col_w[ci] * PT, h * PT, fill=1, stroke=1)
                        pdf.setFillColor(HexColor(PAPER if ri == 0 else INK))
                        pdf.setFont(fonts[('sans', ri == 0)], size)
                        for li, t in enumerate(ls):
                            pdf.drawString((xx + 0.12) * PT, Y(yy + 0.1 + (li * 1.25 + 0.92) * size / 72), t)
                        xx += col_w[ci]
                    yy += h
        pdf.showPage()
    pdf.save()
    return buf.getvalue()


# ------------------------------------------------------------------ darslikka qo'shish
AUTO_NOTE = 'Darslik matnidan avtomatik tayyorlangan'


def make_for_topic(topic):
    """Mavzu uchun PDF + PPTX taqdimot (avvalgi avtomatik nusxalar almashtiriladi, o'qituvchi yuklaganlari qoladi)."""
    from django.core.files.base import ContentFile

    from .models import Resource

    for old in topic.resources.filter(kind='slides', note__startswith=AUTO_NOTE):
        old.file.delete(save=False)
        old.delete()
    n = len(outline(topic))
    base = f'{topic.course.slug}-{topic.number:02d}-mavzu'
    pdf = Resource(course=topic.course, topic=topic, kind='slides', title=f'{topic.number}-mavzu taqdimoti',
                   note=f'{AUTO_NOTE} · {n} ta slayd · brauzerda ko‘rish uchun')
    pdf.file.save(f'{base}.pdf', ContentFile(render_pdf(topic)), save=True)
    ppt = Resource(course=topic.course, topic=topic, kind='slides', title=f'{topic.number}-mavzu taqdimoti (PowerPoint)',
                   note=f'{AUTO_NOTE} · darsda ko‘rsatish va tahrirlash uchun')
    ppt.file.save(f'{base}.pptx', ContentFile(render_pptx(topic)), save=True)
    return n
